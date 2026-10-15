"""Multi-format document ingestion.

Walks a directory, dispatches each supported file to a format-specific loader,
normalises whitespace, computes a SHA-256 content hash for idempotent
re-ingestion, and persists rows to the ``documents`` table.

Supported formats: ``.pdf`` (pdfplumber), ``.docx`` (python-docx),
``.eml`` (stdlib ``email``), ``.txt`` (plain text).
"""

from __future__ import annotations

import email
import email.policy
import hashlib
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pdfplumber
from sqlalchemy import select
from sqlalchemy.orm import Session

from .db import Document, get_database

# --------------------------------------------------------------------------
# Loaders
# --------------------------------------------------------------------------


def load_pdf(path: str | Path) -> str:
    """Extract text from a PDF using pdfplumber."""
    text_parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


def load_docx(path: str | Path) -> str:
    """Extract paragraph + table text from a DOCX using python-docx."""
    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_eml(path: str | Path) -> str:
    """Extract headers + plain-text body from an EML file (stdlib email)."""
    raw = Path(path).read_bytes()
    msg = email.message_from_bytes(raw, policy=email.policy.default)
    header_lines = []
    for header in ("From", "To", "Date", "Subject"):
        value = msg.get(header)
        if value:
            header_lines.append(f"{header}: {value}")
    body = ""
    body_part = msg.get_body(preferencelist=("plain",))
    if body_part is not None:
        body = body_part.get_content()
    elif not msg.is_multipart():
        body = msg.get_content()
    return "\n".join(header_lines) + "\n\n" + body


def load_txt(path: str | Path) -> str:
    """Read a plain-text file."""
    return Path(path).read_text(encoding="utf-8", errors="replace")


LOADERS: dict[str, Callable[[str | Path], str]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".eml": load_eml,
    ".txt": load_txt,
}

DOC_TYPES: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".eml": "eml",
    ".txt": "txt",
}

SUPPORTED_SUFFIXES = tuple(LOADERS)


def load_file(path: str | Path) -> tuple[str, str]:
    """Dispatch a single file to its loader. Returns ``(pristine_text, doc_type)``."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in LOADERS:
        raise ValueError(f"Unsupported document type: {suffix!r}")
    return LOADERS[suffix](p), DOC_TYPES[suffix]


# --------------------------------------------------------------------------
# Normalisation + hashing
# --------------------------------------------------------------------------

_WS_RE = re.compile(r"[ \t\f\v]+")
_MULTI_NL_RE = re.compile(r"\n{3,}")


def normalize_text(text: str) -> str:
    """Collapse runs of whitespace while preserving paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = _MULTI_NL_RE.sub("\n\n", text)
    return text.strip()


def sha256_text(text: str) -> str:
    """SHA-256 of the normalised text, used for idempotent re-ingest."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------
# Directory ingestion
# --------------------------------------------------------------------------


@dataclass
class IngestReport:
    """Outcome of a directory ingest."""

    ingested: list[Document] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)
    skipped: list[tuple[str, str]] = field(default_factory=list)

    @property
    def new_count(self) -> int:
        return len(self.ingested)

    @property
    def duplicate_count(self) -> int:
        return len(self.duplicates)

    @property
    def skipped_count(self) -> int:
        return len(self.skipped)

    # Convenience so a report can be treated as a list of ingested Documents.
    def __iter__(self) -> Iterator[Document]:
        return iter(self.ingested)

    def __len__(self) -> int:
        return len(self.ingested)

    def __getitem__(self, index: int) -> Document:
        return self.ingested[index]

    def summary(self) -> str:
        return (
            f"{self.new_count} ingested / {self.duplicate_count} duplicates "
            f"/ {self.skipped_count} skipped"
        )


def iter_files(root: str | Path) -> Iterator[Path]:
    """Yield supported files under ``root`` in deterministic (sorted) order."""
    root = Path(root)
    if root.is_file():
        yield root
        return
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in LOADERS:
            yield path


def ingest_file(path: Path, session: Session) -> Document | None:
    """Ingest one file. Returns the new Document, or None if it is a duplicate."""
    text, doc_type = load_file(path)
    normalized = normalize_text(text)
    digest = sha256_text(normalized)
    existing = session.scalar(select(Document).where(Document.sha256 == digest))
    if existing is not None:
        return None
    document = Document(
        source_path=str(path),
        doc_type=doc_type,
        raw_text=normalized,
        char_count=len(normalized),
        sha256=digest,
    )
    session.add(document)
    return document


def ingest_dir(
    path: str | Path,
    session: Session | None = None,
    database_url: str | None = None,
) -> IngestReport:
    """Walk ``path`` and ingest every supported document.

    Duplicate content (identical SHA-256) already present in the database is
    skipped, making re-ingestion idempotent. Returns an :class:`IngestReport`
    which behaves like a list of the newly ingested :class:`Document` rows.
    """
    report = IngestReport()
    own_session = session is None
    db = get_database(database_url)
    if own_session:
        session = db.session()
    assert session is not None
    try:
        for file_path in iter_files(path):
            try:
                document = ingest_file(file_path, session)
            except Exception as exc:  # pragma: no cover - corrupt file path
                report.skipped.append((str(file_path), str(exc)))
                continue
            if document is None:
                report.duplicates.append(str(file_path))
            else:
                report.ingested.append(document)
        if own_session:
            session.commit()
            for document in report.ingested:
                session.refresh(document)
    finally:
        if own_session:
            session.close()
    return report
